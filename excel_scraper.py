from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import xlwings as xw


def main():
    wb = xw.Book.caller()
    sht = wb.sheets["Settings"]

    # Read the table into a DataFrame
    tbl = sht.tables["tbl_SETTINGS"].range.options(pd.DataFrame, index=False).value

    # Get the input folder
    input_folder = sht.range("INPUT_FOLDER").value
    input_path = Path(input_folder)

    # Prepare an empty list to collect results
    results = []

    # Get the list of Excel files
    excel_files = list(input_path.glob("*.xlsx"))

    # Create the output folder if it doesn't exist
    output_folder = Path(wb.fullname).parent / "output"
    output_folder.mkdir(parents=True, exist_ok=True)

    # Iterate over all Excel files in the input folder
    for excel_file in excel_files:
        wb_file = openpyxl.load_workbook(excel_file, data_only=True)

        # Iterate over the rows in the table
        for _, row in tbl.iterrows():
            output_name = row["value_name"]
            sheet_name = row["sheet_name"]
            cell_address = row["cell"]

            # Read the value
            if sheet_name in wb_file.sheetnames:
                sheet = wb_file[sheet_name]
                value = sheet[cell_address].value

            # Append the result to the results list
            results.append([str(excel_file.name), output_name, value])

    # Create a DataFrame from the results
    df_results = pd.DataFrame(results, columns=["File Name", "Value Name", "Value"])

    # Create the output file path with a timestamp
    output_file = output_folder / f'results_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

    # Save the results DataFrame to an Excel file
    df_results.to_excel(output_file, index=False)

    # Notify the user
    wb.app.alert(f"Results saved to {output_file}", "Success")


if __name__ == "__main__":
    xw.Book("excel_scraper.xlsm").set_mock_caller()
    main()
